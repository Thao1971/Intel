"""Normalize broad regional sector-multiple workbooks for Intel."""
from __future__ import annotations
from pathlib import Path
from statistics import median
from typing import Any, Dict, Iterable, List, Mapping, Optional

BENCHMARK_VERSION = "regional-sector-benchmarks-v1"
METRIC_FIELDS = {
    "per": "PER trailing",
    "per_forward": "PER forward",
    "price_to_book": "P/VC",
    "ev_revenue": "VE/ventas",
    "ev_ebitda": "VE/EBITDA positivo",
    "ev_ebitda_all": "VE/EBITDA todas",
    "ev_ebit": "VE/EBIT",
}
GROUP_ARCHETYPES = {
    "Construcción e ingeniería": "construction",
    "Consumo": "retail",
    "Energía": "energy_utilities",
    "Financiero": "financial_services",
    "Inmobiliario": "real_estate",
    "Medios y publicidad": "media_content",
    "Salud": "healthcare",
    "Software y datos": "software_data",
    "Turismo y ocio": "hospitality",
    "Otros": "general_business",
}
# Exact industries used when a broad Arroba group contains distinct business models.
INDUSTRY_ARCHETYPES = {
    "Advertising": "advertising",
    "Broadcasting": "media_content",
    "Entertainment": "media_content",
    "Publishing & Newspapers": "media_content",
    "Social Media": "media_content",
    "Computer Services": "software_data",
    "Software (Entertainment)": "software_data",
    "Software (Internet)": "software_data",
    "Software (System & Application)": "software_data",
    "Engineering/Construction": "construction",
    "Construction Supplies": "construction",
    "Real Estate (Development)": "real_estate",
    "Real Estate (General/Diversified)": "real_estate",
    "Real Estate (Operations & Services)": "real_estate",
    "R.E.I.T.": "real_estate",
    "Restaurants/Dining": "hospitality",
    "Hotel/Gaming": "hospitality",
    "Air Transport": "transport_logistics",
    "Transportation": "transport_logistics",
}

def _number(value: Any) -> Optional[float]:
    if isinstance(value, bool) or value in (None, ""):
        return None
    try:
        value=float(value)
        return value if value == value else None
    except (TypeError, ValueError):
        return None

def _header_map(values: Iterable[Any]) -> Dict[str, int]:
    return {str(value).strip(): index for index,value in enumerate(values) if value is not None}

def load_sector_benchmark(path: Path, region: str, as_of: str,
                          provider: str = "damodaran") -> Dict[str, Any]:
    from openpyxl import load_workbook
    workbook=load_workbook(path,read_only=True,data_only=True)
    sheet=workbook["Sectores"]
    iterator=sheet.iter_rows(values_only=True)
    headers=_header_map(next(iterator))
    required={"Industria Damodaran","Grupo Arroba","Empresas",*METRIC_FIELDS.values()}
    missing=sorted(required-set(headers))
    if missing:
        raise ValueError(f"Missing benchmark columns: {', '.join(missing)}")
    records=[]
    for values in iterator:
        industry=values[headers["Industria Damodaran"]]
        if not industry:
            continue
        group=str(values[headers["Grupo Arroba"]] or "")
        industry=str(industry).strip()
        metrics={key:_number(values[headers[column]]) for key,column in METRIC_FIELDS.items()}
        records.append({
            "provider":provider,
            "benchmark_version":BENCHMARK_VERSION,
            "as_of":as_of,
            "region":region,
            "industry":industry,
            "group_arroba":group,
            "archetype":INDUSTRY_ARCHETYPES.get(industry, GROUP_ARCHETYPES.get(group,"general_business")),
            "company_count":int(values[headers["Empresas"]] or 0),
            "metrics":metrics,
            "source_file":path.name,
            "source_sheet":"Sectores",
            "active":True,
        })
    return {"provider":provider,"region":region,"as_of":as_of,
            "benchmark_version":BENCHMARK_VERSION,"source_file":path.name,
            "record_count":len(records),"records":records}

def _percentile(values: List[float], probability: float) -> Optional[float]:
    if not values: return None
    ordered=sorted(values); position=(len(ordered)-1)*probability
    lower=int(position); upper=min(lower+1,len(ordered)-1); fraction=position-lower
    return ordered[lower]+(ordered[upper]-ordered[lower])*fraction

def select_benchmark(records: Iterable[Mapping[str,Any]], archetype: str,
                     metric: str) -> Optional[Dict[str,Any]]:
    usable=[record for record in records
            if record.get("archetype")==archetype
            and _number((record.get("metrics") or {}).get(metric)) not in (None,0)]
    if not usable and archetype=="advertising":
        usable=[record for record in records
                if record.get("industry")=="Advertising"
                and _number((record.get("metrics") or {}).get(metric)) not in (None,0)]
    if not usable:
        return None
    values=[float(record["metrics"][metric]) for record in usable]
    return {
        "metric":metric,
        "multiple":median(values),
        "p25":_percentile(values,.25),"median":median(values),"p75":_percentile(values,.75),
        "industries":[record["industry"] for record in usable],
        "industry_count":len(usable),
        "company_count":sum(int(record.get("company_count") or 0) for record in usable),
        "regions":sorted({str(record.get("region")) for record in usable}),
        "providers":sorted({str(record.get("provider")) for record in usable}),
        "as_of":max(str(record.get("as_of") or "") for record in usable),
        "status":"observed_sector_benchmark",
        "benchmark_version":BENCHMARK_VERSION,
    }
